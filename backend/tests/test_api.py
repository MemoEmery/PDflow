"""Rodar: cd backend && pip install pytest httpx && pytest -q
(testes de LibreOffice exigem `soffice` instalado; são pulados se não houver)"""
import io, json, os, shutil, sys, tempfile, time
from pathlib import Path
os.environ["STORAGE_DIR"] = tempfile.mkdtemp(); os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ.pop("REDIS_URL", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pymupdf as fitz, pytest
from PIL import Image
from fastapi.testclient import TestClient
from main import app

app.state.limiter.enabled = False  # os testes fazem dezenas de conversões por minuto
c = TestClient(app)

def run(tool, files, opts=None):
    r = c.post("/api/convert", data={"tool": tool, "options": json.dumps(opts or {})}, files=files)
    if r.status_code != 202: return r.status_code, r.json()["detail"], None
    jid = r.json()["job_id"]
    for _ in range(100):
        j = c.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("completed", "failed"): break
        time.sleep(.2)
    return j["status"], j.get("filename") or j.get("error"), jid

def pdf(n=3):
    d = fitz.open()
    for i in range(n): d.new_page().insert_text((72, 72), f"P{i+1}")
    return d.tobytes()

def pages(jid): return [p.get_text().strip() for p in fitz.open(stream=c.get(f"/api/files/{jid}").content)]
F = lambda b, n="a.pdf", m="application/pdf": [("files", (n, b, m))]

def test_txt_to_pdf():
    st, name, jid = run("to-pdf", F(b"ola", "a.txt", "text/plain"))
    assert (st, name) == ("completed", "a.pdf") and "ola" in pages(jid)[0]

@pytest.mark.skipif(not shutil.which("soffice"), reason="sem LibreOffice")
def test_csv_to_pdf_sem_prefixo_no_cabecalho():
    st, _, jid = run("to-pdf", F(b"a,b\n1,2\n", "dados.csv", "text/csv"))
    assert st == "completed" and "000_" not in pages(jid)[0]

def test_rejeita_formato_e_pdf_falso():
    assert run("to-pdf", F(b"x", "a.exe", "application/octet-stream"))[0] == 400
    assert run("merge-pdf", F(b"nao pdf"))[0] == 400
    assert run("tool-inexistente", F(pdf()))[0] == 400

def test_ferramentas_de_paginas():
    assert pages(run("extract-pages", F(pdf(5)), {"pages": "2-3,5"})[2]) == ["P2", "P3", "P5"]
    assert pages(run("remove-pages", F(pdf(5)), {"pages": "1,4-5"})[2]) == ["P2", "P3"]
    assert pages(run("reorder-pages", F(pdf(3)), {"pages": "3,1"})[2]) == ["P3", "P1", "P2"]
    assert run("remove-pages", F(pdf(2)), {"pages": "1-2"})[0] == "failed"
    assert run("extract-pages", F(pdf(2)), {"pages": "9"})[0] == "failed"

def test_proteger_e_desbloquear():
    jid = run("protect-pdf", F(pdf()), {"password": "abcd"})[2]
    locked = c.get(f"/api/files/{jid}").content
    assert fitz.open(stream=locked).needs_pass
    assert run("unlock-pdf", F(locked), {"password": "errada"})[0] == "failed"
    assert pages(run("unlock-pdf", F(locked), {"password": "abcd"})[2]) == ["P1", "P2", "P3"]

def test_imagens_opcoes():
    b = io.BytesIO(); Image.new("RGB", (800, 400), "red").save(b, "PNG")
    jid = run("to-pdf", F(b.getvalue(), "i.png", "image/png"), {"size": "A4", "margin": "medium"})[2]
    assert [round(x) for x in fitz.open(stream=c.get(f"/api/files/{jid}").content)[0].rect[2:]] == [842, 595]
    assert run("to-pdf", F(b.getvalue(), "i.png", "image/png"), {"size": "A3"})[0] == "failed"


def test_converter_de_pdf():
    jid = run("pdf-to-txt", F(pdf(2)))[2]
    assert "P1" in c.get(f"/api/files/{jid}").text
    st, name, jid = run("pdf-to-excel", F(pdf(2)))
    assert (st, name) == ("completed", "a.xlsx")
    import openpyxl; wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/files/{jid}").content))
    assert [r[1] for r in wb["Texto"].iter_rows(min_row=2, values_only=True)] == ["P1", "P2"]
    st, name, jid = run("pdf-to-pptx", F(pdf(3)))
    from pptx import Presentation
    assert (st, name) == ("completed", "a.pptx") and len(Presentation(io.BytesIO(c.get(f"/api/files/{jid}").content)).slides) == 3
    assert run("pdf-to-word", F(pdf(1)))[:2] == ("completed", "a.docx")
    blank = fitz.open(); blank.new_page()
    assert run("pdf-to-txt", F(blank.tobytes()))[0] == "failed"  # sem texto (digitalizado)


def test_juntar_pdfs_na_ordem_enviada():
    def doc(tag, n):
        d = fitz.open()
        for i in range(n): d.new_page().insert_text((72, 72), f"{tag}{i+1}")
        return d.tobytes()
    st, name, jid = run("merge-pdf", [("files", ("b.pdf", doc("B", 2), "application/pdf")), ("files", ("a.pdf", doc("A", 3), "application/pdf"))])
    assert (st, name) == ("completed", "juntado.pdf")
    assert pages(jid) == ["B1", "B2", "A1", "A2", "A3"]          # respeita a ordem enviada
    assert run("merge-pdf", F(doc("A", 1)))[0] == 400            # um só arquivo não basta
    assert run("merge-pdf", [("files", ("a.pdf", doc("A", 1), "application/pdf")), ("files", ("x.pdf", b"nao pdf", "application/pdf"))])[0] == 400
