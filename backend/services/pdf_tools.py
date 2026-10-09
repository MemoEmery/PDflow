import re, secrets, zipfile
from pathlib import Path
import pymupdf as fitz  # PyMuPDF

PDF_TOOLS = {"merge-pdf", "split-pdf", "compress-pdf", "pdf-to-jpg", "pdf-to-png", "pdf-to-word", "pdf-to-txt", "pdf-to-excel", "pdf-to-pptx",
             "extract-pages", "remove-pages", "reorder-pages", "rotate-pdf", "watermark-pdf", "page-numbers-pdf", "protect-pdf", "unlock-pdf"}
MAX_PAGES = 200
NO_TEXT = "Este PDF não tem texto selecionável (parece ser digitalizado)."

def _open(p: Path, allow_locked=False):
    d = fitz.open(p)
    if d.needs_pass and not allow_locked: raise ValueError("Este PDF está protegido por senha. Use a ferramenta Desbloquear PDF.")
    if not d.needs_pass and d.page_count > MAX_PAGES: raise ValueError(f"O PDF excede o limite de {MAX_PAGES} páginas.")
    return d

def _txt(opts: dict, key: str, default: str = "", maxlen: int = 100) -> str:
    v = opts.get(key, default)
    if not isinstance(v, str): raise ValueError("Opção inválida.")
    return v.strip()[:maxlen]

def parse_pages(spec: str, n: int) -> list[int]:
    """'1-3,5,8-' -> índices 0-based, na ordem informada."""
    out: list[int] = []
    tokens = [t.strip() for t in spec.split(",") if t.strip()]
    if not tokens or len(tokens) > 200: raise ValueError("Informe as páginas, por exemplo: 1-3,5.")
    for t in tokens:
        m = re.fullmatch(r"(\d{1,5})(?:(-)(\d{0,5}))?", t)
        if not m: raise ValueError("Intervalo de páginas inválido. Use o formato 1-3,5.")
        a = int(m[1]); b = (int(m[3]) if m[3] else n) if m[2] else a
        if not (1 <= a <= n and 1 <= b <= n): raise ValueError(f"As páginas devem estar entre 1 e {n}.")
        out += range(a - 1, b) if a <= b else range(a - 1, b - 2, -1)
    return out

def run_pdf_tool(tool: str, files: list[Path], work: Path, opts: dict | None = None) -> Path:
    opts = opts or {}; f = files[0]; base = re.sub(r"^\d{3}_", "", f.stem)
    if tool == "merge-pdf":
        out = fitz.open()
        for p in files:
            with _open(p) as d: out.insert_pdf(d)
        o = work / "juntado.pdf"; out.save(str(o), garbage=3, deflate=True); return o
    if tool == "compress-pdf":
        o = work / f"{base}_comprimido.pdf"
        with _open(f) as d:
            if hasattr(d, "rewrite_images"): d.rewrite_images(dpi_threshold=150, dpi_target=110, quality=70)
            d.save(str(o), garbage=4, deflate=True, clean=True)
        return o
    if tool == "pdf-to-word":
        from pdf2docx import Converter
        _open(f).close(); o = work / f"{base}.docx"; cv = Converter(str(f))
        try: cv.convert(str(o))
        finally: cv.close()
        return o
    if tool in {"extract-pages", "remove-pages", "reorder-pages"}:
        o = work / f"{base}_{tool.split('-')[0]}.pdf"
        with _open(f) as d:
            sel = parse_pages(_txt(opts, "pages", "", 400), d.page_count)
            if tool == "remove-pages":
                gone = set(sel); sel = [i for i in range(d.page_count) if i not in gone]
                if not sel: raise ValueError("Não é possível remover todas as páginas.")
            elif tool == "reorder-pages":
                sel += [i for i in range(d.page_count) if i not in set(sel)]  # o que não foi citado vai ao final
            d.select(sel); d.save(str(o), garbage=3, deflate=True)
        return o
    if tool == "pdf-to-txt":
        with _open(f) as d: text = "\n\n".join(pg.get_text().strip() for pg in d).strip()
        if not text: raise ValueError(NO_TEXT)
        o = work / f"{base}.txt"; o.write_text(text, encoding="utf-8"); return o
    if tool == "pdf-to-excel":
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active); found = False
        with _open(f) as d:
            for i, pg in enumerate(d, 1):
                for k, tb in enumerate(pg.find_tables().tables, 1):
                    rows = tb.extract()
                    if not rows: continue
                    found = True; ws = wb.create_sheet(f"Pág{i}-Tab{k}"[:31])
                    for row in rows: ws.append([("" if c is None else str(c).replace("\n", " ")) for c in row])
            if not found:  # sem tabelas: uma linha da planilha por linha de texto
                lines = [(i, ln) for i, pg in enumerate(d, 1) for ln in pg.get_text().splitlines() if ln.strip()]
                if not lines: raise ValueError(NO_TEXT)
                ws = wb.create_sheet("Texto"); ws.append(["Página", "Texto"])
                for i, ln in lines: ws.append([i, ln])
        o = work / f"{base}.xlsx"; wb.save(o); return o
    if tool == "pdf-to-pptx":
        from pptx import Presentation
        from pptx.util import Emu
        import io
        prs = Presentation(); W = 12192000
        with _open(f) as d:
            r0 = d[0].rect; prs.slide_width, prs.slide_height = Emu(W), Emu(int(W * r0.height / r0.width))
            for pg in d:  # cada página vira um slide com a imagem da página (não editável)
                sl = prs.slides.add_slide(prs.slide_layouts[6])
                sl.shapes.add_picture(io.BytesIO(pg.get_pixmap(dpi=110).tobytes("png")), 0, 0, prs.slide_width, prs.slide_height)
        o = work / f"{base}.pptx"; prs.save(o); return o
    if tool == "rotate-pdf":
        ang = _txt(opts, "angle", "90")
        if ang not in {"90", "180", "270"}: raise ValueError("Ângulo inválido.")
        o = work / f"{base}_girado.pdf"
        with _open(f) as d:
            for pg in d: pg.set_rotation((pg.rotation + int(ang)) % 360)
            d.save(str(o), garbage=3, deflate=True)
        return o
    if tool == "watermark-pdf":
        text = _txt(opts, "text", "CONFIDENCIAL", 60) or "CONFIDENCIAL"
        o = work / f"{base}_marca.pdf"
        with _open(f) as d:
            for pg in d:
                r = pg.rect; size = max(24, min(r.width, r.height) / 9)
                w = fitz.get_text_length(text, fontname="helv", fontsize=size)
                c = fitz.Point(r.width / 2, r.height / 2)
                pg.insert_text(fitz.Point(c.x - w / 2, c.y), text, fontname="helv", fontsize=size,
                               color=(0.5, 0.5, 0.5), fill_opacity=0.3, morph=(c, fitz.Matrix(45)), overlay=True)
            d.save(str(o), garbage=3, deflate=True)
        return o
    if tool == "page-numbers-pdf":
        fmt = _txt(opts, "format", "n"); o = work / f"{base}_numerado.pdf"
        with _open(f) as d:
            total = d.page_count
            for i, pg in enumerate(d):
                label = f"{i+1} / {total}" if fmt == "n/N" else str(i + 1)
                w = fitz.get_text_length(label, fontname="helv", fontsize=10)
                pg.insert_text(fitz.Point((pg.rect.width - w) / 2, pg.rect.height - 24), label, fontname="helv", fontsize=10)
            d.save(str(o), garbage=3, deflate=True)
        return o
    if tool == "protect-pdf":
        pw = _txt(opts, "password", "", 128)
        if len(pw) < 4: raise ValueError("Use uma senha com pelo menos 4 caracteres.")
        o = work / f"{base}_protegido.pdf"
        with _open(f) as d:
            d.save(str(o), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=pw, owner_pw=secrets.token_hex(16), garbage=3, deflate=True)
        return o
    if tool == "unlock-pdf":
        pw = _txt(opts, "password", "", 128); o = work / f"{base}_desbloqueado.pdf"
        with _open(f, allow_locked=True) as d:
            if d.needs_pass and not d.authenticate(pw): raise ValueError("Senha incorreta.")
            d.save(str(o), encryption=fitz.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
        return o
    o = work / (f"{base}_paginas.zip" if tool == "split-pdf" else f"{base}_{tool[-3:]}.zip")
    with zipfile.ZipFile(o, "w", zipfile.ZIP_DEFLATED) as z, _open(f) as d:
        for i, page in enumerate(d):
            n = f"pagina_{i+1:03d}"
            if tool == "split-pdf":
                part = fitz.open(); part.insert_pdf(d, from_page=i, to_page=i)
                z.writestr(n + ".pdf", part.tobytes(deflate=True)); part.close()
            else:
                ext = "jpg" if tool == "pdf-to-jpg" else "png"
                z.writestr(f"{n}.{ext}", page.get_pixmap(dpi=150).tobytes("jpeg" if ext == "jpg" else "png"))
    return o
